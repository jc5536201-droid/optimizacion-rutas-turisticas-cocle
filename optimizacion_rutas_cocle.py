"""
OPTIMIZACIÓN DE RUTAS TURÍSTICAS – PROVINCIA DE COCLÉ, PANAMÁ
Algoritmo de Dijkstra | 29 nodos | 53 aristas | Menú interactivo
Universidad de Panamá – Facultad de Ciencias Naturales y Exactas Licenciatura en Matemáticas

7 DÍAS · 4-6 LUGARES POR DÍA · DATOS ESTIMADOS DESDE GOOGLE MAPS
CRITERIO OFICIAL DEL PROYECTO: DISTANCIA (según indicación del asesor)

QUÉ HACE EL PROGRAMA
  - Dijkstra (implementación propia, heap) calcula los caminos mínimos entre
    TODOS los pares de nodos, para distancia, tiempo y costo.
  - El ORDEN de visita de cada día se resuelve con búsqueda exhaustiva
    (permutaciones, problema del viajante/TSP) sobre esas distancias mínimas,
    saliendo y regresando a Penonomé (nodo 15). Dijkstra NO ordena visitas.
  - La AGRUPACIÓN de lugares por día (DIAS_CONFIG) es manual, por zonas.

VALIDACIONES (Sección 3.6.1)
  a) Validación externa REAL: se genera una plantilla CSV con una muestra de
     aristas; tú anotas valores de una fuente independiente (OSRM/OpenRoute-
     Service/Google Maps en otra fecha) y el programa calcula los errores.
  b) Consistencia interna: subestructura óptima, TODOS los pares de nodos.
  c) Sensibilidad: (i) ¿cambia la ruta si se optimiza por tiempo en vez de
     distancia?  (ii) Monte Carlo: se perturban las distancias ±10% / ±20% y
     se mide qué tan estable es el orden óptimo de cada día.

PRUEBAS COMPUTACIONALES (Sección 3.6.3)
  a) Caja negra: manual (opción 5) y automática con tabla de casos (opción 9-2)
  b) Integridad de datos: automática y real (opción 9-1)
  c) Tiempo de ejecución (Tabla 9) (opción 9-3)

MODO SIN MENÚ (genera todas las tablas de una vez):
     python rutas_cocle_dijkstra.py --pruebas

NOTA SOBRE EL COSTO:
  costo = distancia * COSTO_POR_KM (constante). Al ser una transformación
  lineal, optimizar por costo da exactamente la misma ruta que por distancia.
  Por eso el criterio OFICIAL es DISTANCIA; tiempo y costo son métricas de
  apoyo. Justifica COSTO_POR_KM en la tesis (consumo del vehículo y precio
  del combustible, con fuente y fecha).
"""

import csv
import heapq
import os
import random
import statistics
import sys
import time
from itertools import permutations
from math import radians, sin, cos, asin, sqrt


# ── Consola: UTF-8 y colores ANSI (evita errores con emojis en Windows) ───────
def _configurar_consola():
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8")
        except Exception:
            pass
    if os.name == "nt":
        os.system("")  # activa secuencias ANSI en Windows 10+


_configurar_consola()


# ── Verificar e instalar dependencias ─────────────────────────────────────────
def verificar_dependencias():
    faltantes = []
    for pkg in ["networkx", "matplotlib", "openpyxl"]:
        try:
            __import__(pkg)
        except ImportError:
            faltantes.append(pkg)
    if faltantes:
        print(f"\n  Instalando dependencias: {', '.join(faltantes)} ...")
        os.system(f'"{sys.executable}" -m pip install {" ".join(faltantes)} -q')
        print("  ¡Listo!\n")


verificar_dependencias()

import networkx as nx
import matplotlib


# ── Selección automática del backend ──────────────────────────────────────────
def _seleccionar_backend():
    if os.name == "nt":
        return "TkAgg"
    if sys.platform == "darwin":
        return "TkAgg"
    if "DISPLAY" in os.environ:
        return "TkAgg"
    return "Agg"


try:
    matplotlib.use(_seleccionar_backend())
    import matplotlib.pyplot as plt
    _fig_test = plt.figure()
    plt.close(_fig_test)
except Exception:
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

import matplotlib.patches as mpatches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ═══════════════════════════════════════════════════════════════════════════════
# 0. CONSTANTES
# ═══════════════════════════════════════════════════════════════════════════════

BASE = 15                      # Penonomé: punto de salida y regreso de cada día
COSTO_POR_KM = 0.15            # USD/km (JUSTIFICAR en la tesis con fuente y fecha)
CRITERIOS = ("distancia", "tiempo", "costo")
TOL = 1e-3                     # tolerancia numérica al comparar matrices

# ═══════════════════════════════════════════════════════════════════════════════
# 1. DATOS DEL SISTEMA TURÍSTICO (29 NODOS)
# ═══════════════════════════════════════════════════════════════════════════════

ATRACTIVOS = {
    # Nodos originales (1-21)
    1:  {"nombre": "Playa Santa Clara",               "cod": "PSC", "tipo": "Playa",            "puntaje": 27, "distrito": "Antón"},
    2:  {"nombre": "Playa Farallón",                  "cod": "PFA", "tipo": "Playa",            "puntaje": 26, "distrito": "Antón"},
    3:  {"nombre": "Playa El Salado",                 "cod": "PES", "tipo": "Playa",            "puntaje": 21, "distrito": "Aguadulce"},
    4:  {"nombre": "Playa Blanca",                    "cod": "PBL", "tipo": "Playa",            "puntaje": 26, "distrito": "Antón"},
    5:  {"nombre": "Playa Juan Hombrón",              "cod": "PJH", "tipo": "Playa",            "puntaje": 20, "distrito": "Antón"},
    6:  {"nombre": "Mercado Artesanía Valle Antón",   "cod": "MAV", "tipo": "Cultural",         "puntaje": 26, "distrito": "Antón"},
    7:  {"nombre": "Serpentario Maravillas Tropicales","cod": "SMT", "tipo": "Naturaleza",       "puntaje": 24, "distrito": "Antón"},
    8:  {"nombre": "Museo Hermanos Arias Madrid",     "cod": "MHA", "tipo": "Cultural/Hist.",   "puntaje": 25, "distrito": "Penonomé"},
    9:  {"nombre": "P.N. Omar Torrijos",              "cod": "PNT", "tipo": "Parque Nacional",  "puntaje": 22, "distrito": "La Pintada"},
    10: {"nombre": "Sitio Arqueológico El Caño",      "cod": "SAC", "tipo": "Arqueológico",     "puntaje": 26, "distrito": "Natá"},
    11: {"nombre": "Museo Regional Stella Sierra",    "cod": "MSS", "tipo": "Cultural/Hist.",   "puntaje": 22, "distrito": "Aguadulce"},
    12: {"nombre": "Iglesia San Juan Bautista",       "cod": "ISJ", "tipo": "Histórico",        "puntaje": 24, "distrito": "Penonomé"},
    13: {"nombre": "El Chorro Las Yayas",             "cod": "CLY", "tipo": "Cascada",          "puntaje": 25, "distrito": "La Pintada"},
    14: {"nombre": "Balneario Las Mendozas",          "cod": "BLM", "tipo": "Balneario",        "puntaje": 21, "distrito": "Penonomé"},
    15: {"nombre": "Penonomé",                        "cod": "PEN", "tipo": "Hub/Ciudad",       "puntaje": 28, "distrito": "Penonomé"},
    16: {"nombre": "Aguadulce",                       "cod": "AGU", "tipo": "Hub/Ciudad",       "puntaje": 25, "distrito": "Aguadulce"},
    17: {"nombre": "Antón",                           "cod": "ANT", "tipo": "Hub/Ciudad",       "puntaje": 23, "distrito": "Antón"},
    18: {"nombre": "La Pintada",                      "cod": "LAP", "tipo": "Hub/Ciudad",       "puntaje": 22, "distrito": "La Pintada"},
    19: {"nombre": "Natá",                            "cod": "NAT", "tipo": "Hub/Ciudad",       "puntaje": 23, "distrito": "Natá"},
    20: {"nombre": "Parroquia Ntra. Sra. Candelaria", "cod": "PNC", "tipo": "Histórico",        "puntaje": 21, "distrito": "La Pintada"},
    21: {"nombre": "Cerro Gaital",                    "cod": "CGA", "tipo": "Montaña",          "puntaje": 27, "distrito": "Antón"},
    # Nuevos nodos (22-29)
    22: {"nombre": "Museo de Penonomé",               "cod": "MPE", "tipo": "Cultural",         "puntaje": 24, "distrito": "Penonomé"},
    23: {"nombre": "Mercado Artesanías La Pintada",   "cod": "MLA", "tipo": "Cultural",         "puntaje": 23, "distrito": "La Pintada"},
    24: {"nombre": "Balneario Los Algarrobos",        "cod": "BAL", "tipo": "Naturaleza",       "puntaje": 22, "distrito": "La Pintada"},
    25: {"nombre": "Iglesia Santiago Apóstol",        "cod": "ISA", "tipo": "Histórico",        "puntaje": 24, "distrito": "Natá"},
    26: {"nombre": "Ecoparque Don Arcelio",           "cod": "ECO", "tipo": "Naturaleza",       "puntaje": 22, "distrito": "Natá"},
    27: {"nombre": "Salinas de Aguadulce",            "cod": "SAL", "tipo": "Naturaleza",       "puntaje": 21, "distrito": "Aguadulce"},
    28: {"nombre": "Mariposario",                     "cod": "MAR", "tipo": "Naturaleza",       "puntaje": 24, "distrito": "Antón"},
    29: {"nombre": "Canopy Adventure",                "cod": "CAN", "tipo": "Aventura",         "puntaje": 25, "distrito": "Antón"},
}

# Coordenadas reales (latitud, longitud) en grados decimales, por ID de nodo.
# OPCIONAL pero recomendado (tu parte GIS): si las llenas, la prueba de
# integridad verifica que ninguna distancia por carretera sea menor que la
# distancia en línea recta. Ejemplo:  15: (8.5190, -80.3570),8
# NO inventes valores: cópialos de tu capa GIS o de Google Maps.
COORDENADAS = {}

# ARISTAS CON DATOS ESTIMADOS DESDE GOOGLE MAPS: (origen, destino, km, minutos)
ARISTAS_RAW = [
    # Penonomé (PEN) → Conexiones originales
    (15, 17, 22.3, 27),   # PEN → ANT
    (15, 10, 30.1, 33),   # PEN → SAC
    (15, 16, 50.2, 49),   # PEN → AGU
    (15, 18, 18.4, 23),   # PEN → LAP
    (15, 8, 4.1, 7),      # PEN → MHA
    (15, 12, 4.6, 9),     # PEN → ISJ
    (15, 14, 4.6, 9),     # PEN → BLM
    (15, 21, 42.9, 78),   # PEN → CGA

    # Antón (ANT) → Conexiones originales
    (17, 1, 23.1, 25),    # ANT → PSC
    (17, 2, 21.4, 25),    # ANT → PFA
    (17, 4, 19.8, 22),    # ANT → PBL
    (17, 5, 18.0, 24),    # ANT → PJH
    (17, 6, 37.6, 59),    # ANT → MAV
    (17, 7, 35.3, 62),    # ANT → SMT
    (17, 21, 40.5, 61),   # ANT → CGA

    # Playas (Conexiones costeras)
    (1, 2, 6.8, 11),      # PSC → PFA
    (2, 4, 8.4, 12),      # PFA → PBL
    (4, 5, 22.0, 24),     # PBL → PJH
    (1, 5, 27.3, 32),     # PSC → PJH

    # Valle de Antón
    (6, 7, 2.2, 5),       # MAV → SMT
    (6, 21, 2.9, 7),      # MAV → CGA

    # Natá (NAT) → Conexiones originales
    (19, 10, 28.7, 37),   # NAT → SAC
    (19, 17, 66.0, 66),   # NAT → ANT
    (19, 16, 19.4, 31),   # NAT → AGU

    # Aguadulce (AGU) → Conexiones originales
    (16, 10, 26.8, 27),   # AGU → SAC
    (16, 11, 2.8, 6),     # AGU → MSS
    (16, 3, 10.1, 16),      # AGU → PES

    # La Pintada (LAP) → Conexiones originales
    (18, 13, 34.2, 54),   # LAP → CLY
    (18, 20, 0.13, 1),    # LAP → PNC
    (18, 9, 41.3, 88),    # LAP → PNT

    # Parque Nacional → Conexiones originales
    (9, 13, 7.1, 34),     # PNT → CLY
    (9, 10, 44.7, 85),    # PNT → SAC

    # NUEVOS NODOS - Penonomé
    (15, 22, 4.3, 8),     # PEN → MPE
    (8, 22, 1.5, 4),      # MHA → MPE
    (22, 12, 0.8, 3),     # MPE → ISJ

    # NUEVOS NODOS - La Pintada
    (18, 23, 0.8, 1),     # LAP → MLA
    (18, 24, 0.35, 2),    # LAP → BAL
    (23, 24, 1.1, 2),     # MLA → BAL
    (20, 23, 0.8, 1),     # PNC → MLA
    (20, 24, 0.3, 1),     # PNC → BAL

    # NUEVOS NODOS - Natá
    (10, 25, 13.8, 15),   # SAC → ISA
    (19, 25, 16.0, 28),   # NAT → ISA
    (25, 26, 11.4, 13),   # ISA → ECO
    (19, 26, 26.3, 35),   # NAT → ECO

    # NUEVOS NODOS - Aguadulce
    (16, 27, 8.7, 15),    # AGU → SAL
    (3, 27, 1.4, 2),      # PES → SAL
    (27, 11, 7.0, 10),    # SAL → MSS

    # NUEVOS NODOS - Antón
    (17, 28, 33.7, 56),   # ANT → MAR
    (17, 29, 39.1, 56),   # ANT → CAN
    (28, 7, 2.1, 5),      # MAR → SMT
    (28, 29, 3.6, 9),    # MAR → CAN
    (29, 21, 5.8, 12),    # CAN → CGA

    # Arista urbana adicional (Día 2)
    (12, 14, 1.0, 3),     # ISJ → BLM
]

COLORES_TIPO = {
    "Hub/Ciudad":      "#1D9E75",
    "Playa":           "#2E86C1",
    "Cultural":        "#F39C12",
    "Cultural/Hist.":  "#F39C12",
    "Histórico":       "#F39C12",
    "Naturaleza":      "#8E44AD",
    "Parque Nacional": "#27AE60",
    "Cascada":         "#2980B9",
    "Balneario":       "#1ABC9C",
    "Arqueológico":    "#C0392B",
    "Montaña":         "#27AE60",
    "Aventura":        "#E67E22",
}

# POSICIONES ESQUEMÁTICAS (NO geográficas, no a escala): solo para dibujar.
POSICIONES = {
    1:  (0.95, 0.2978),   2:  (0.9087, 0.1582), 3:  (0.0669, 0.05),
    4:  (0.7873, 0.0871), 5:  (0.6609, 0.1479), 6:  (0.6452, 0.6373),
    7:  (0.7461, 0.7708), 8:  (0.3431, 0.5761), 9:  (0.0714, 0.6373),
    10: (0.3238, 0.2871), 11: (0.1875, 0.259),  12: (0.5911, 0.4386),
    13: (0.0852, 0.8172), 14: (0.5051, 0.6505), 15: (0.4694, 0.5092),
    16: (0.1923, 0.1129), 17: (0.7707, 0.3921), 18: (0.2507, 0.8295),
    19: (0.1856, 0.4052), 20: (0.1473, 0.948),  21: (0.6452, 0.9015),
    22: (0.3525, 0.4302), 23: (0.3294, 0.95),   24: (0.2079, 0.6904),
    25: (0.0517, 0.2278), 26: (0.05, 0.437),    27: (0.3311, 0.1054),
    28: (0.8643, 0.8476), 29: (0.8272, 0.6521),
}

# ═══════════════════════════════════════════════════════════════════════════════
# 2. RUTAS DE 7 DÍAS
#    distancia_km / tiempo_auto_min / costo_usd / orden_optimo se calculan en
#    vivo en main() con recalcular_estadisticas_dias().
# ═══════════════════════════════════════════════════════════════════════════════

DIAS_CONFIG = {
    1: {"destinos": [1, 2, 4, 5, 17],
        "zona": "🌊 Playas de Antón", "color": "#2E86C1",
        "descripcion": "Recorrido por las playas más hermosas de la costa de Antón",
        "frase": "🌅 El sol, la arena y el mar te esperan", "nivel": "Fácil"},
    2: {"destinos": [8, 22, 12, 14, 15],
        "zona": "🏛️ Penonomé Histórico", "color": "#F39C12",
        "descripcion": "Recorrido urbano por los tesoros históricos de Penonomé",
        "frase": "🏛️ Descubre la historia de Coclé en cada esquina", "nivel": "Fácil"},
    3: {"destinos": [18, 20, 23, 24, 13, 9],
        "zona": "⛰️ La Pintada - Montaña y Cultura", "color": "#27AE60",
        "descripcion": "Aventura en la montaña: parroquia, mercado, balneario, cascada y parque",
        "frase": "⛰️ La montaña te llama con sus cascadas y paisajes", "nivel": "Alta"},
    4: {"destinos": [10, 25, 26, 19],
        "zona": "🏺 Ruta Arqueológica de Natá", "color": "#C0392B",
        "descripcion": "Viaje al pasado: El Caño, iglesia colonial, ecoparque y ciudad",
        "frase": "🗿 Donde la historia precolombina cobra vida", "nivel": "Media"},
    5: {"destinos": [6, 7, 28, 21, 29],
        "zona": "🌿 Naturaleza de Antón", "color": "#8E44AD",
        "descripcion": "Un día en el valle: mercado, serpentario, mariposario, cerro y canopy",
        "frase": "🦋 Donde la naturaleza y la aventura se encuentran", "nivel": "Media"},
    6: {"destinos": [16, 3, 27, 11],
        "zona": "🌅 Tesoros de Aguadulce", "color": "#1ABC9C",
        "descripcion": "Un día completo: playa, salinas, museo y ciudad",
        "frase": "🧂 Donde el mar y la historia se fusionan", "nivel": "Fácil"},
    7: {"destinos": [15, 18, 20, 23, 24],
        "zona": "🎯 Circuito Integrador", "color": "#E67E22",
        "descripcion": "Un recorrido por lo mejor de Penonomé y La Pintada",
        "frase": "🎯 Un viaje que une ciudades, historia y naturaleza", "nivel": "Media"},
}
for _cfg in DIAS_CONFIG.values():
    _cfg.update({"distancia_km": 0.0, "tiempo_auto_min": 0, "costo_usd": 0.0, "orden_optimo": None})

# ═══════════════════════════════════════════════════════════════════════════════
# 3. ALGORITMO DE DIJKSTRA
# ═══════════════════════════════════════════════════════════════════════════════

def es_id_valido(x):
    return isinstance(x, int) and x in ATRACTIVOS


def construir_grafo():
    G = nx.Graph()
    for nid, data in ATRACTIVOS.items():
        G.add_node(nid, **data)
    for u, v, dist, tiempo in ARISTAS_RAW:
        # Sin redondeo por arista: el costo es exactamente distancia*tarifa.
        G.add_edge(u, v, distancia=dist, tiempo=tiempo, costo=dist * COSTO_POR_KM)
    return G


def dijkstra(G, origen, criterio):
    INF = float('inf')
    dist = {n: INF for n in G.nodes()}
    prev = {n: None for n in G.nodes()}
    dist[origen] = 0
    heap = [(0, origen)]
    visitados = set()
    while heap:
        d_u, u = heapq.heappop(heap)
        if u in visitados:
            continue
        visitados.add(u)
        for v in G.neighbors(u):
            alt = dist[u] + G[u][v][criterio]
            if alt < dist[v]:
                dist[v] = alt
                prev[v] = u
                heapq.heappush(heap, (alt, v))
    return dist, prev


def reconstruir_camino(prev, origen, destino):
    camino = []
    actual = destino
    while actual is not None:
        camino.append(actual)
        actual = prev[actual]
    camino.reverse()
    return camino if camino and camino[0] == origen else []


def calcular_todas_matrices(G):
    nodos = sorted(G.nodes())
    matrices, caminos = {}, {}
    for criterio in CRITERIOS:
        mat, cam = {}, {}
        for origen in nodos:
            dist_min, prev = dijkstra(G, origen, criterio)
            mat[origen] = {d: (round(dist_min[d], 4) if dist_min[d] != float('inf') else None) for d in nodos}
            cam[origen] = {d: reconstruir_camino(prev, origen, d) for d in nodos}
        matrices[criterio] = mat
        caminos[criterio] = cam
    return matrices, caminos


def matriz_un_criterio(G, criterio):
    """Matriz de distancias mínimas sin redondeo (para Monte Carlo)."""
    return {o: dijkstra(G, o, criterio)[0] for o in G.nodes()}


def ruta_optima_dia(destinos, matrices, criterio="distancia"):
    """TSP por fuerza bruta sobre distancias mínimas (Dijkstra), base = Penonomé."""
    destinos = [d for d in destinos if d != BASE]
    mat = matrices[criterio]
    mejor_costo = float('inf')
    mejor_orden = None
    for perm in permutations(destinos):
        secuencia = [BASE] + list(perm) + [BASE]
        total = 0
        valida = True
        for i in range(len(secuencia) - 1):
            c = mat[secuencia[i]][secuencia[i + 1]]
            if c is None:
                valida = False
                break
            total += c
        if valida and total < mejor_costo:
            mejor_costo = total
            mejor_orden = list(perm)
    return mejor_orden, mejor_costo


def recalcular_estadisticas_dias(matrices, criterio="distancia"):
    """Recalcula distancia/tiempo/costo y guarda el ORDEN óptimo de cada día."""
    for dia, cfg in DIAS_CONFIG.items():
        destinos = [d for d in cfg["destinos"] if d != BASE]
        orden, _ = ruta_optima_dia(destinos, matrices, criterio)
        if orden is None:
            print(f"  ⚠ No se pudo recalcular el Día {dia} (ruta inválida).")
            continue
        sec = [BASE] + orden + [BASE]
        pares = list(zip(sec[:-1], sec[1:]))
        cfg["orden_optimo"] = orden
        cfg["distancia_km"] = round(sum(matrices["distancia"][a][b] for a, b in pares), 1)
        cfg["tiempo_auto_min"] = round(sum(matrices["tiempo"][a][b] for a, b in pares))
        cfg["costo_usd"] = round(sum(matrices["costo"][a][b] for a, b in pares), 2)


def texto_ruta_dia(cfg):
    orden = cfg["orden_optimo"] if cfg["orden_optimo"] else [d for d in cfg["destinos"] if d != BASE]
    sec = [BASE] + orden + [BASE]
    return " → ".join(ATRACTIVOS[n]["cod"] for n in sec)

# ═══════════════════════════════════════════════════════════════════════════════
# 4. FUNCIONES DE VISUALIZACIÓN
# ═══════════════════════════════════════════════════════════════════════════════

_CONTADOR_GRAFICO = [0]


def _mostrar(fig=None, nombre=None):
    if fig is None:
        fig = plt.gcf()
    _CONTADOR_GRAFICO[0] += 1
    if matplotlib.get_backend().lower() == "agg":
        archivo = nombre or f"grafo_{_CONTADOR_GRAFICO[0]}.png"
        fig.savefig(archivo, dpi=150, bbox_inches="tight", facecolor=fig.get_facecolor())
        print(f"  ℹ Sin pantalla gráfica: imagen guardada como {archivo}")
    else:
        plt.show()


def dibujar_grafo_completo(G, criterio="distancia", guardar=None):
    fig, ax = plt.subplots(figsize=(15, 11))
    ax.set_facecolor("#F0F4F8")
    fig.patch.set_facecolor("#E8EDF2")

    for u, v in G.edges():
        x0, y0 = POSICIONES[u]
        x1, y1 = POSICIONES[v]
        ax.plot([x0, x1], [y0, y1], color="#BBBBBB", linewidth=1.0, zorder=1, alpha=0.6)

    RADIO_HUB, RADIO_NORM = 0.024, 0.017
    for nid in G.nodes():
        x, y = POSICIONES[nid]
        tipo = G.nodes[nid]["tipo"]
        color = COLORES_TIPO.get(tipo, "#999999")
        es_hub = tipo == "Hub/Ciudad"
        radio = RADIO_HUB if es_hub else RADIO_NORM
        ax.add_patch(plt.Circle((x + 0.003, y - 0.003), radio, color="#00000015", zorder=3))
        ax.add_patch(plt.Circle((x, y), radio, color=color, zorder=4, linewidth=1.5, ec="white"))
        ax.text(x, y, G.nodes[nid]["cod"], ha="center", va="center",
                fontsize=8.5 if es_hub else 7, fontweight="bold", color="white", zorder=5)

    nodos_con_nombre = [15, 17, 16, 18, 19, 1, 2, 4, 21, 13, 10, 9, 20]
    for nid in nodos_con_nombre:
        x, y = POSICIONES[nid]
        nombre = G.nodes[nid]["nombre"]
        if len(nombre) > 22:
            nombre = nombre[:20] + "..."
        if x > 0.78:
            lx, ly, ha = x + 0.035, y, "left"
        elif x < 0.22:
            lx, ly, ha = x - 0.035, y, "right"
        elif y > 0.82:
            lx, ly, ha = x, y + 0.038, "center"
        else:
            lx, ly, ha = x, y - 0.038, "center"
        ax.text(lx, ly, nombre, ha=ha, va="center", fontsize=7, fontweight="500",
                color="#3A3A3A", zorder=6,
                bbox=dict(boxstyle="round,pad=0.18", facecolor="white",
                          edgecolor="#DDDDDD", alpha=0.9, linewidth=0.4))

    leyenda = [
        mpatches.Patch(color="#1D9E75", label="Hub/Ciudad"),
        mpatches.Patch(color="#2E86C1", label="Playa"),
        mpatches.Patch(color="#F39C12", label="Cultural/Hist."),
        mpatches.Patch(color="#8E44AD", label="Naturaleza"),
        mpatches.Patch(color="#C0392B", label="Arqueológico"),
        mpatches.Patch(color="#27AE60", label="Parque/Montaña"),
        mpatches.Patch(color="#E67E22", label="Aventura"),
    ]
    ax.legend(handles=leyenda, loc="lower left", fontsize=8.5, framealpha=0.95,
              edgecolor="#AAAAAA", fancybox=True, ncol=2)
    ax.set_title(f"Grafo turístico - Provincia de Coclé, Panamá (diagrama esquemático, no a escala)\n"
                 f"{len(G.nodes())} nodos · {len(G.edges())} aristas",
                 fontsize=13, fontweight="bold", color="#1F4E79", pad=14)
    ax.set_xlim(-0.06, 1.06)
    ax.set_ylim(-0.06, 1.06)
    ax.axis("off")
    ax.set_aspect("equal")
    plt.tight_layout(pad=1.5)

    if guardar:
        plt.savefig(guardar, dpi=150, bbox_inches="tight", facecolor=fig.get_facecolor())
        print(f"  ✔ Grafo guardado: {guardar}")
    else:
        _mostrar(fig, f"grafo_completo_{criterio}.png")
    plt.close(fig)


def dibujar_grafo_dia(G, dia, matrices, caminos, criterio="distancia", guardar=None):
    cfg = DIAS_CONFIG[dia]
    destinos = [d for d in cfg["destinos"] if d != BASE]
    color_dia = cfg["color"]

    orden, costo_total = ruta_optima_dia(destinos, matrices, criterio)
    if orden is None:
        print(f"  ✗ No se pudo calcular la ruta para el Día {dia}.")
        return

    secuencia = [BASE] + orden + [BASE]
    nodos_ruta = set(secuencia)
    segmentos_dir = []
    for i in range(len(secuencia) - 1):
        seg = caminos[criterio][secuencia[i]][secuencia[i + 1]]
        for j in range(len(seg) - 1):
            segmentos_dir.append((seg[j], seg[j + 1]))
        nodos_ruta.update(seg)

    pos_sub = {n: POSICIONES[n] for n in nodos_ruta}
    unidad = {"distancia": "km", "tiempo": "min", "costo": "USD"}[criterio]

    fig, ax = plt.subplots(figsize=(12, 9.5))
    ax.set_facecolor("#F0F4F8")
    fig.patch.set_facecolor("#E8EDF2")

    for u, v in G.edges():
        if u in nodos_ruta and v in nodos_ruta:
            x0, y0 = pos_sub[u]
            x1, y1 = pos_sub[v]
            ax.plot([x0, x1], [y0, y1], color="#DDDDDD", linewidth=0.7, alpha=0.5, zorder=1)

    def _mas_cercano(px, py):
        return min(((px - ox) ** 2 + (py - oy) ** 2) ** 0.5 for ox, oy in pos_sub.values())

    ya_etiquetado = set()
    RADIO_ETIQ_LIBRE = 0.052
    for u, v in segmentos_dir:
        x0, y0 = pos_sub[u]
        x1, y1 = pos_sub[v]
        ax.annotate("", xy=(x1, y1), xytext=(x0, y0),
                    arrowprops=dict(arrowstyle="-|>", color=color_dia, lw=2.2,
                                    connectionstyle="arc3,rad=0.05", mutation_scale=17),
                    zorder=3)
        par = tuple(sorted([u, v]))
        if par not in ya_etiquetado and G.has_edge(u, v):
            val = G[u][v][criterio]
            txt = f"{val:.2f} {unidad}" if criterio == "costo" else f"{val:g} {unidad}"
            mx, my = (x0 + x1) / 2, (y0 + y1) / 2
            dx, dy = x1 - x0, y1 - y0
            largo = (dx ** 2 + dy ** 2) ** 0.5 or 1e-6
            perp_x, perp_y = -dy / largo, dx / largo
            offset = 0.040 if largo > 0.09 else 0.062
            lx, ly = mx + perp_x * offset, my + perp_y * offset
            if _mas_cercano(lx, ly) < RADIO_ETIQ_LIBRE:
                encontrado = False
                for signo in (1, -1):
                    for extra in (0.030, 0.055, 0.080):
                        lx2 = mx + perp_x * signo * (offset + extra)
                        ly2 = my + perp_y * signo * (offset + extra)
                        if _mas_cercano(lx2, ly2) >= RADIO_ETIQ_LIBRE:
                            lx, ly = lx2, ly2
                            encontrado = True
                            break
                    if encontrado:
                        break
            ax.text(lx, ly, txt, ha="center", va="center", fontsize=7.5,
                    fontweight="bold", color=color_dia, zorder=5,
                    bbox=dict(boxstyle="round,pad=0.15", facecolor="white",
                              edgecolor=color_dia, alpha=0.95, linewidth=0.6))
            ya_etiquetado.add(par)

    RADIO_BASE, RADIO_DEST, RADIO_TRANS = 0.036, 0.028, 0.018
    nodos_destino = set(destinos)
    nodos_transito = nodos_ruta - nodos_destino - {BASE}
    orden_visita = {}
    for i, nodo in enumerate(secuencia):
        if nodo != BASE and nodo not in orden_visita:
            orden_visita[nodo] = i

    for nid in nodos_ruta:
        x, y = pos_sub[nid]
        if nid == BASE:
            cn, radio, ec, lw = "#1D9E75", RADIO_BASE, "white", 2.2
        elif nid in nodos_destino:
            cn = COLORES_TIPO.get(G.nodes[nid]["tipo"], "#2E86C1")
            radio, ec, lw = RADIO_DEST, color_dia, 2.0
        else:
            cn, radio, ec, lw = "#AAAAAA", RADIO_TRANS, "#888888", 1.2
        ax.add_patch(plt.Circle((x + 0.003, y - 0.003), radio, color="#00000015", zorder=3))
        ax.add_patch(plt.Circle((x, y), radio, color=cn, ec=ec, linewidth=lw, zorder=4))
        cod = G.nodes[nid]["cod"] + ("*" if nid in nodos_transito else "")
        if nid in orden_visita:
            cod = f"{orden_visita[nid]}·{cod}"
        fs_cod = 9 if nid == BASE else (7.5 if nid in nodos_destino else 6.5)
        ax.text(x, y, cod, ha="center", va="center", fontsize=fs_cod,
                fontweight="bold", color="white", zorder=5)

    for nid in nodos_ruta:
        if nid not in nodos_destino and nid != BASE:
            continue
        x, y = pos_sub[nid]
        nombre = G.nodes[nid]["nombre"]
        if len(nombre) > 22:
            nombre = nombre[:20] + "..."
        if x > 0.72:
            lx, ly, ha = x + 0.045, y, "left"
        elif x < 0.28:
            lx, ly, ha = x - 0.045, y, "right"
        elif y > 0.75:
            lx, ly, ha = x, y + 0.05, "center"
        else:
            lx, ly, ha = x, y - 0.05, "center"
        ax.text(lx, ly, nombre, ha=ha, va="center", fontsize=7.5, fontweight="500",
                color="#333333", zorder=6,
                bbox=dict(boxstyle="round,pad=0.15", facecolor="white",
                          edgecolor=color_dia if nid in nodos_destino else "#CCCCCC",
                          alpha=0.92, linewidth=0.5))

    ruta_str = " → ".join(G.nodes[n]["cod"] for n in secuencia)
    ax.text(0.01, 0.01, f"Ruta óptima ({criterio})\n{ruta_str}\nTotal: {costo_total:.1f} {unidad}",
            transform=ax.transAxes, fontsize=8.5, color="#1A1A1A", va="bottom",
            bbox=dict(boxstyle="round,pad=0.35", facecolor="white",
                      edgecolor=color_dia, alpha=0.96, linewidth=1.1))

    leyenda = [
        mpatches.Patch(color="#1D9E75", label="Base: Penonomé"),
        mpatches.Patch(color=color_dia, label="Ruta óptima"),
        mpatches.Patch(color="#AAAAAA", label="Tránsito (*)"),
    ]
    ax.legend(handles=leyenda, loc="lower right", fontsize=8.5, framealpha=0.95,
              edgecolor="#AAAAAA", fancybox=True)

    zona_limpia = cfg["zona"].split(" ", 1)[1] if " " in cfg["zona"] else cfg["zona"]
    ax.set_title(f"DÍA {dia} — {zona_limpia}\nCriterio: {criterio.upper()} | {len(destinos)} destinos (diagrama esquemático)",
                 fontsize=13, fontweight="bold", color=color_dia, pad=14)

    xs = [p[0] for p in pos_sub.values()]
    ys = [p[1] for p in pos_sub.values()]
    margen = 0.14
    ax.set_xlim(min(xs) - margen, max(xs) + margen)
    ax.set_ylim(min(ys) - margen, max(ys) + margen)
    ax.axis("off")
    ax.set_aspect("equal")
    plt.tight_layout(pad=1.5)

    if guardar:
        plt.savefig(guardar, dpi=150, bbox_inches="tight", facecolor=fig.get_facecolor())
        print(f"  ✔ Grafo Día {dia} guardado: {guardar}")
    else:
        _mostrar(fig, f"Grafo_Dia_{dia}_{criterio}.png")
    plt.close(fig)

# ═══════════════════════════════════════════════════════════════════════════════
# 5. PRUEBAS Y VALIDACIONES REALES
# ═══════════════════════════════════════════════════════════════════════════════

def _haversine_km(a, b):
    (lat1, lon1), (lat2, lon2) = a, b
    p1, p2 = radians(lat1), radians(lat2)
    h = sin((p2 - p1) / 2) ** 2 + cos(p1) * cos(p2) * sin(radians(lon2 - lon1) / 2) ** 2
    return 2 * 6371.0088 * asin(sqrt(h))


def _cod(n):
    return ATRACTIVOS[n]["cod"]


# ── 5.1 Integridad de datos (3.6.3.b) ─────────────────────────────────────────
def prueba_integridad_datos(G, verbose=True):
    """Devuelve (errores, advertencias). Errores = datos inválidos.
    Advertencias = datos sospechosos que debes revisar/remedir."""
    errores, advertencias = [], []
    ids = set(ATRACTIVOS)

    # Nodos
    if G.number_of_nodes() != len(ATRACTIVOS):
        errores.append(f"El grafo tiene {G.number_of_nodes()} nodos y el inventario {len(ATRACTIVOS)}.")
    for nid in ATRACTIVOS:
        if nid not in POSICIONES:
            errores.append(f"Nodo {nid} sin posición de dibujo.")
        if ATRACTIVOS[nid]["tipo"] not in COLORES_TIPO:
            errores.append(f"Nodo {nid}: tipo '{ATRACTIVOS[nid]['tipo']}' sin color definido.")
    codigos = [d["cod"] for d in ATRACTIVOS.values()]
    if len(set(codigos)) != len(codigos):
        errores.append("Hay códigos de nodo repetidos.")

    # Aristas
    vistas = set()
    for u, v, d, t in ARISTAS_RAW:
        if u not in ids or v not in ids:
            errores.append(f"Arista ({u},{v}) usa un ID inexistente.")
            continue
        if u == v:
            errores.append(f"Arista ({_cod(u)}) es un autolazo.")
        clave = frozenset((u, v))
        if clave in vistas:
            errores.append(f"Arista duplicada {_cod(u)}-{_cod(v)}.")
        vistas.add(clave)
        if d <= 0 or t <= 0:
            errores.append(f"Arista {_cod(u)}-{_cod(v)} con peso no positivo (d={d}, t={t}).")
    if len(vistas) != G.number_of_edges():
        errores.append(f"Aristas únicas en datos ({len(vistas)}) ≠ aristas del grafo ({G.number_of_edges()}).")

    # Conectividad
    if not nx.is_connected(G):
        comps = [sorted(_cod(n) for n in c) for c in nx.connected_components(G)]
        errores.append(f"El grafo NO es conexo. Componentes: {comps}")

    # Días
    cubiertos = set()
    for dia, cfg in DIAS_CONFIG.items():
        dest = cfg["destinos"]
        if any(d not in ids for d in dest):
            errores.append(f"Día {dia}: hay IDs de destino inexistentes.")
            continue
        if len(set(dest)) != len(dest):
            errores.append(f"Día {dia}: destinos repetidos.")
        if not 4 <= len(dest) <= 6:
            errores.append(f"Día {dia}: tiene {len(dest)} lugares (se esperan 4 a 6).")
        if BASE in dest:
            advertencias.append(f"Día {dia}: incluye a Penonomé (base) como destino; solo hay "
                                f"{len(dest) - 1} paradas distintas a la base.")
        cubiertos.update(dest)
    sin_dia = sorted(_cod(n) for n in ids - cubiertos)
    if sin_dia:
        advertencias.append(f"Atractivos que no aparecen en ningún día: {sin_dia}")

    # Velocidad implícita
    for u, v, data in G.edges(data=True):
        if data["distancia"] >= 2:
            vel = data["distancia"] / (data["tiempo"] / 60)
            if vel < 10 or vel > 110:
                advertencias.append(f"Arista {_cod(u)}-{_cod(v)}: velocidad implícita {vel:.0f} km/h (¿dato correcto?)")

    # Aristas dominadas
    for u, v, data in G.edges(data=True):
        for crit in ("distancia", "tiempo"):
            H = G.copy()
            H.remove_edge(u, v)
            try:
                alt = nx.dijkstra_path_length(H, u, v, weight=crit)
                camino = nx.dijkstra_path(H, u, v, weight=crit)
            except nx.NetworkXNoPath:
                continue
            if alt < data[crit] - 1e-9:
                advertencias.append(
                    f"Arista {_cod(u)}-{_cod(v)} ({crit}={data[crit]:g}) es MÁS LARGA que el camino alterno "
                    f"{' → '.join(_cod(n) for n in camino)} ({alt:g}). Vuelve a medirla.")

    # Coordenadas (opcional)
    if COORDENADAS:
        faltan = sorted(_cod(n) for n in ids if n not in COORDENADAS)
        if faltan:
            advertencias.append(f"COORDENADAS incompleto; faltan: {faltan}")
        for u, v, data in G.edges(data=True):
            if u in COORDENADAS and v in COORDENADAS:
                h = _haversine_km(COORDENADAS[u], COORDENADAS[v])
                if h > data["distancia"] + 0.05:
                    errores.append(f"Arista {_cod(u)}-{_cod(v)}: carretera {data['distancia']} km < línea recta {h:.2f} km (imposible).")
                elif h > 0.5 and data["distancia"] / h > 3:
                    advertencias.append(f"Arista {_cod(u)}-{_cod(v)}: carretera/línea recta = {data['distancia'] / h:.1f} (muy sinuosa).")

    if verbose:
        print("\n" + "=" * 70)
        print("  PRUEBA DE INTEGRIDAD DE DATOS (3.6.3.b)")
        print("=" * 70)
        print(f"  Nodos: {G.number_of_nodes()} | Aristas: {G.number_of_edges()} | Conexo: {nx.is_connected(G)}")
        print(f"  Coordenadas para prueba de línea recta: {'sí' if COORDENADAS else 'NO (COORDENADAS vacío, prueba omitida)'}")
        if errores:
            print(f"\n  ❌ ERRORES ({len(errores)}):")
            for e in errores:
                print(f"    - {e}")
        else:
            print("\n  ✅ Sin errores de integridad.")
        if advertencias:
            print(f"\n  ⚠️  ADVERTENCIAS ({len(advertencias)}) – revisar/remedir, no son fallos del código:")
            for a in advertencias:
                print(f"    - {a}")
        print("=" * 70 + "\n")
    return errores, advertencias


# ── 5.2 Verificación del algoritmo contra implementaciones independientes ────
def verificacion_algoritmo(G, matrices, caminos, verbose=True):
    """Compara TODAS las matrices contra networkx (Dijkstra) y Floyd-Warshall,
    valida cada camino reconstruido, simetría, diagonal e inecuación triangular."""
    nodos = sorted(G.nodes())
    fallos, chequeos = [], 0
    for crit in CRITERIOS:
        ref = dict(nx.all_pairs_dijkstra_path_length(G, weight=crit))
        fw = nx.floyd_warshall(G, weight=crit)
        mat, cam = matrices[crit], caminos[crit]
        for u in nodos:
            if abs(mat[u][u]) > TOL:
                fallos.append(f"[{crit}] diagonal {_cod(u)} ≠ 0")
            for v in nodos:
                chequeos += 1
                m = mat[u][v]
                if m is None or abs(m - ref[u][v]) > TOL or abs(m - fw[u][v]) > TOL:
                    fallos.append(f"[{crit}] {_cod(u)}→{_cod(v)}: propio={m}, networkx={ref[u][v]:.4f}, FW={fw[u][v]:.4f}")
                    continue
                if abs(m - mat[v][u]) > TOL:
                    fallos.append(f"[{crit}] asimetría {_cod(u)}↔{_cod(v)}")
                p = cam[u][v]
                if not p or p[0] != u or p[-1] != v:
                    fallos.append(f"[{crit}] camino mal formado {_cod(u)}→{_cod(v)}")
                    continue
                if any(not G.has_edge(p[i], p[i + 1]) for i in range(len(p) - 1)):
                    fallos.append(f"[{crit}] camino {_cod(u)}→{_cod(v)} usa una arista inexistente")
                    continue
                suma = sum(G[p[i]][p[i + 1]][crit] for i in range(len(p) - 1))
                if abs(suma - m) > TOL:
                    fallos.append(f"[{crit}] suma del camino {_cod(u)}→{_cod(v)} ({suma:.4f}) ≠ matriz ({m})")
        for u in nodos:
            for v in nodos:
                for w in nodos:
                    chequeos += 1
                    if mat[u][v] > mat[u][w] + mat[w][v] + TOL:
                        fallos.append(f"[{crit}] triangular violada {_cod(u)}→{_cod(v)} vía {_cod(w)}")
    if verbose:
        print("\n" + "=" * 70)
        print("  VERIFICACIÓN DEL ALGORITMO vs networkx y Floyd-Warshall (3 criterios)")
        print("=" * 70)
        print(f"  Comprobaciones realizadas: {chequeos}")
        if fallos:
            print(f"  ❌ Fallos: {len(fallos)}")
            for f in fallos[:10]:
                print(f"    - {f}")
        else:
            print("  ✅ Todas las matrices y caminos coinciden con las implementaciones independientes.")
        print("=" * 70 + "\n")
    return len(fallos) == 0, chequeos


def tabla_casos_caja_negra(G, matrices, caminos, archivo="caja_negra_resultados.csv"):
    """Tabla de casos de prueba documentada (3.6.3.a), criterio distancia."""
    nodos = sorted(G.nodes())
    lejano = max(((u, v) for u in nodos for v in nodos if u < v),
                 key=lambda p: matrices["distancia"][p[0]][p[1]])
    pares = [(15, 8), (15, 17), (18, 20), (15, 21), (1, 7), (19, 26), (18, 9),
             (3, 11), (5, 29), (20, 24), lejano]
    filas = []
    print("\n" + "=" * 96)
    print("  PRUEBA DE CAJA NEGRA AUTOMÁTICA (3.6.3.a) – criterio DISTANCIA (km)")
    print("=" * 96)
    print(f"  {'#':<3} {'Caso':<22} {'Esperado':>9} {'Obtenido':>9} {'Resultado':>10}  Ruta obtenida")
    print("-" * 96)
    n = 0
    for u, v in pares:
        n += 1
        esperado = nx.dijkstra_path_length(G, u, v, weight="distancia")
        obtenido = matrices["distancia"][u][v]
        ok = abs(esperado - obtenido) <= TOL
        ruta = " → ".join(_cod(x) for x in caminos["distancia"][u][v])
        caso = f"{_cod(u)} → {_cod(v)}" + (" (más lejano)" if (u, v) == lejano else "")
        print(f"  {n:<3} {caso:<22} {esperado:>9.2f} {obtenido:>9.2f} {'✅ OK' if ok else '❌ FALLA':>10}  {ruta}")
        filas.append([n, caso, round(esperado, 4), round(obtenido, 4), "OK" if ok else "FALLA", ruta])
    # Casos límite
    n += 1
    ok = abs(matrices["distancia"][15][15]) <= TOL
    print(f"  {n:<3} {'origen = destino':<22} {0:>9.2f} {matrices['distancia'][15][15]:>9.2f} {'✅ OK' if ok else '❌ FALLA':>10}  PEN")
    filas.append([n, "origen = destino", 0, matrices["distancia"][15][15], "OK" if ok else "FALLA", "PEN"])
    for malo in (0, 30):
        n += 1
        ok = not es_id_valido(malo)
        print(f"  {n:<3} {'ID inválido ' + str(malo):<22} {'rechazar':>9} {'rechaza' if ok else 'acepta':>9} {'✅ OK' if ok else '❌ FALLA':>10}  -")
        filas.append([n, f"ID inválido {malo}", "rechazar", "rechaza" if ok else "acepta", "OK" if ok else "FALLA", "-"])
    aprobados = sum(1 for f in filas if f[4] == "OK")
    print("-" * 96)
    print(f"  RESULTADO: {aprobados}/{len(filas)} casos aprobados")
    print("=" * 96 + "\n")
    with open(archivo, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["caso_n", "caso", "esperado_km", "obtenido_km", "resultado", "ruta"])
        w.writerows(filas)
    print(f"  ✔ Tabla guardada: {archivo}")
    return aprobados == len(filas)


# ── 5.3 Validación interna (3.6.1.b) ─────────────────────────────────────────
def validar_consistencia_interna(G, matrices, caminos, verbose=True):
    """Subestructura óptima en TODOS los pares: todo subcamino de un camino
    mínimo debe costar exactamente la distancia mínima entre sus extremos."""
    nodos = sorted(G.nodes())
    errores, validaciones = [], 0
    for crit in CRITERIOS:
        mat, cam = matrices[crit], caminos[crit]
        for o in nodos:
            for d in nodos:
                if o == d:
                    continue
                p = cam[o][d]
                w = [G[p[i]][p[i + 1]][crit] for i in range(len(p) - 1)]
                for i in range(len(p) - 1):
                    for j in range(i + 1, len(p)):
                        validaciones += 1
                        if abs(sum(w[i:j]) - mat[p[i]][p[j]]) > TOL:
                            errores.append((crit, p[i], p[j]))
    if verbose:
        print("\n" + "=" * 70)
        print("  b) CONSISTENCIA INTERNA (subestructura óptima, todos los pares)")
        print("=" * 70)
        print(f"  Subcaminos verificados: {validaciones}")
        if errores:
            print(f"  ❌ Errores: {len(errores)}")
            for c, a, b in errores[:5]:
                print(f"    - [{c}] subcamino {_cod(a)}→{_cod(b)} no es óptimo")
        else:
            print("  ✅ Todos los subcaminos son óptimos.")
        print("=" * 70 + "\n")
    return len(errores) == 0, validaciones


# ── 5.4 Sensibilidad (3.6.1.c) ───────────────────────────────────────────────
def analisis_sensibilidad(G, matrices, caminos, iteraciones=300,
                          perturbaciones=(0.10, 0.20), semilla=2026, verbose=True):
    nodos = sorted(G.nodes())
    pares = [(u, v) for u in nodos for v in nodos if u < v]
    difieren = [(u, v) for u, v in pares
                if caminos["distancia"][u][v] != caminos["tiempo"][u][v]]
    if verbose:
        print("\n" + "=" * 70)
        print("  c) ANÁLISIS DE SENSIBILIDAD")
        print("=" * 70)
        print("  (i) ¿Cambia el camino mínimo si se optimiza por TIEMPO en vez de DISTANCIA?")
        print(f"      {len(difieren)} de {len(pares)} pares ({len(difieren) / len(pares) * 100:.1f}%) tienen un camino distinto.")
        for u, v in difieren[:5]:
            print(f"      {_cod(u)}→{_cod(v)}: dist {' → '.join(_cod(x) for x in caminos['distancia'][u][v])}"
                  f"  |  tiempo {' → '.join(_cod(x) for x in caminos['tiempo'][u][v])}")
        print("      (costo = distancia × tarifa, produce siempre la misma ruta que distancia)")

    base = {}
    for dia, cfg in DIAS_CONFIG.items():
        dest = [d for d in cfg["destinos"] if d != BASE]
        base[dia] = ruta_optima_dia(dest, matrices, "distancia")

    rng = random.Random(semilla)
    H = G.copy()
    resultados = {}
    for p in perturbaciones:
        estable = {d: 0 for d in DIAS_CONFIG}
        totales = {d: [] for d in DIAS_CONFIG}
        regret = {d: [] for d in DIAS_CONFIG}
        for _ in range(iteraciones):
            for u, v in H.edges():
                H[u][v]["distancia"] = G[u][v]["distancia"] * (1 + rng.uniform(-p, p))
            mat = {"distancia": matriz_un_criterio(H, "distancia")}
            for dia, cfg in DIAS_CONFIG.items():
                dest = [d for d in cfg["destinos"] if d != BASE]
                o, c = ruta_optima_dia(dest, mat, "distancia")
                bo = base[dia][0]
                if o == bo or o == bo[::-1]:
                    estable[dia] += 1
                sec = [BASE] + bo + [BASE]
                costo_base = sum(mat["distancia"][a][b] for a, b in zip(sec[:-1], sec[1:]))
                regret[dia].append((costo_base - c) / c * 100 if c > 0 else 0.0)
                totales[dia].append(c)
        resultados[p] = {d: (estable[d] / iteraciones * 100, statistics.mean(regret[d]), max(regret[d]),
                             min(totales[d]), statistics.mean(totales[d]), max(totales[d]))
                         for d in DIAS_CONFIG}
        if verbose:
            print(f"\n  (ii) Monte Carlo: distancias perturbadas ±{p * 100:.0f}% | {iteraciones} corridas (semilla {semilla})")
            print(f"      {'Día':<4} {'Mismo orden':>11} {'Sobrecosto medio':>17} {'Sobrecosto máx':>15} {'Dist. base':>11} {'Mín':>7} {'Máx':>7}")
            for d in DIAS_CONFIG:
                est, rm, rx, mn, me, mx = resultados[p][d]
                print(f"      {d:<4} {est:>10.1f}% {rm:>16.2f}% {rx:>14.2f}% {base[d][1]:>11.1f} {mn:>7.1f} {mx:>7.1f}")
    if verbose:
        print("\n  Lectura: 'Mismo orden' = % de corridas en que el orden óptimo no cambia.")
        print("  'Sobrecosto' = % extra de distancia si se mantiene el orden base cuando las distancias")
        print("  reales difieren. Un orden puede cambiar con frecuencia y aun así costar casi lo mismo")
        print("  (p. ej. lugares a <1 km entre sí), por eso el sobrecosto es la medida más relevante.")
        print("=" * 70 + "\n")
    return {"pares_distintos": len(difieren), "pares_total": len(pares), "montecarlo": resultados}


# ── 5.5 Validación externa REAL (3.6.1.a) ────────────────────────────────────
COLUMNAS_VALIDACION = ["origen", "destino", "distancia_modelo_km", "tiempo_modelo_min",
                       "distancia_ref_km", "tiempo_ref_min", "fuente", "fecha"]


def generar_plantilla_validacion(G, n=20, archivo="validacion_externa_plantilla.csv", semilla=2026):
    aristas = sorted(tuple(sorted(e)) for e in G.edges())
    muestra = sorted(random.Random(semilla).sample(aristas, min(n, len(aristas))))
    with open(archivo, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUMNAS_VALIDACION)
        for u, v in muestra:
            w.writerow([_cod(u), _cod(v), G[u][v]["distancia"], G[u][v]["tiempo"], "", "", "", ""])
    print(f"\n  ✔ Plantilla creada: {archivo} ({len(muestra)} aristas, muestra aleatoria reproducible)")
    print("    Completa distancia_ref_km, tiempo_ref_min, fuente y fecha con una fuente INDEPENDIENTE")
    print("    (OSRM, OpenRouteService, o Google Maps medido otro día/hora). Luego usa 'Evaluar'.")


def _num(s):
    s = (s or "").strip().replace(",", ".")
    return float(s) if s else None


def evaluar_validacion_externa(G, archivo="validacion_externa_plantilla.csv",
                               tol_rel=0.05, tol_abs_km=0.1, tol_abs_min=1.0, umbral=95.0):
    """Regla: una medida aprueba si |modelo−ref| ≤ max(5% de ref, tolerancia absoluta).
    La tolerancia absoluta (0.1 km / 1 min) evita fallos injustos en tramos muy cortos."""
    if not os.path.exists(archivo):
        print(f"  ✗ No existe '{archivo}'. Genera la plantilla primero.")
        return None
    with open(archivo, newline="", encoding="utf-8-sig") as f:
        texto = f.read()
    primera = texto.splitlines()[0] if texto else ""
    delim = ";" if primera.count(";") > primera.count(",") else ","
    filas = list(csv.DictReader(texto.splitlines(), delimiter=delim))
    cod2id = {d["cod"]: n for n, d in ATRACTIVOS.items()}
    res = []
    print("\n" + "=" * 92)
    print(f"  a) VALIDACIÓN EXTERNA (tolerancia ±{tol_rel * 100:.0f}% o ±{tol_abs_km} km / ±{tol_abs_min:g} min)")
    print("=" * 92)
    print(f"  {'Tramo':<10} {'Dist modelo':>11} {'Dist ref':>9} {'Err %':>7} {'Tiempo mod':>11} {'Tiempo ref':>10} {'Err %':>7} {'Estado':>8}")
    print("-" * 92)
    for r in filas:
        try:
            dref, tref = _num(r.get("distancia_ref_km")), _num(r.get("tiempo_ref_min"))
        except ValueError:
            print(f"  ⚠ Fila {r.get('origen')}-{r.get('destino')}: valor no numérico, se omite.")
            continue
        if dref is None or tref is None:
            continue
        u, v = cod2id.get((r.get("origen") or "").strip().upper()), cod2id.get((r.get("destino") or "").strip().upper())
        if u is None or v is None or not G.has_edge(u, v):
            print(f"  ⚠ Fila {r.get('origen')}-{r.get('destino')}: no corresponde a una arista del grafo, se omite.")
            continue
        dm, tm = G[u][v]["distancia"], G[u][v]["tiempo"]
        ed = (dm - dref) / dref * 100 if dref > 0 else 0.0
        et = (tm - tref) / tref * 100 if tref > 0 else 0.0
        ok_d = abs(dm - dref) <= max(tol_rel * dref, tol_abs_km) + 1e-9
        ok_t = abs(tm - tref) <= max(tol_rel * tref, tol_abs_min) + 1e-9
        ok = ok_d and ok_t
        print(f"  {_cod(u)}-{_cod(v):<6} {dm:>11.2f} {dref:>9.2f} {ed:>+7.1f} {tm:>11.0f} {tref:>10.0f} {et:>+7.1f} {'✅' if ok else '❌':>8}")
        res.append({"tramo": f"{_cod(u)}-{_cod(v)}", "dm": dm, "dref": dref, "ed": ed, "tm": tm, "tref": tref,
                    "et": et, "ok": ok, "fuente": r.get("fuente", ""), "fecha": r.get("fecha", "")})
    if not res:
        print("  ✗ La plantilla no tiene filas completadas (distancia_ref_km y tiempo_ref_min).")
        print("=" * 92 + "\n")
        return None
    apro = sum(1 for r in res if r["ok"])
    pct = apro / len(res) * 100
    mape_d = statistics.mean(abs(r["ed"]) for r in res)
    mape_t = statistics.mean(abs(r["et"]) for r in res)
    print("-" * 92)
    print(f"  Aristas evaluadas: {len(res)} | Aprobadas: {apro} ({pct:.1f}%)")
    print(f"  Error porcentual absoluto medio (MAPE): distancia {mape_d:.2f}% | tiempo {mape_t:.2f}%")
    print(f"  Error máximo: distancia {max(abs(r['ed']) for r in res):.1f}% | tiempo {max(abs(r['et']) for r in res):.1f}%")
    print(f"  {'✅ CUMPLE' if pct >= umbral else '❌ NO CUMPLE'} el umbral de {umbral:.0f}% de aprobación")
    print("=" * 92 + "\n")
    salida = "validacion_externa_resultados.csv"
    with open(salida, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["tramo", "dist_modelo", "dist_ref", "error_dist_%", "tiempo_modelo", "tiempo_ref",
                    "error_tiempo_%", "aprobado", "fuente", "fecha"])
        for r in res:
            w.writerow([r["tramo"], r["dm"], r["dref"], round(r["ed"], 2), r["tm"], r["tref"],
                        round(r["et"], 2), "SI" if r["ok"] else "NO", r["fuente"], r["fecha"]])
    print(f"  ✔ Resultados guardados: {salida}")
    return res

# ═══════════════════════════════════════════════════════════════════════════════
# 6. PRUEBAS DE TIEMPO DE EJECUCIÓN (Tabla 9 – 3.6.3.c)
# ═══════════════════════════════════════════════════════════════════════════════

def probar_tiempo_ejecucion(G, iteraciones=10, semilla=123):
    """Tiempo de calcular Dijkstra desde TODOS los nodos (todos los pares)."""
    rng = random.Random(semilla)
    n_r, m_r = G.number_of_nodes(), G.number_of_edges()
    tamanos = [(5, 7), (10, 15), (21, 32), (n_r, m_r), (50, 120), (100, 300), (200, 600)]
    filas = []
    for nodos, aristas in tamanos:
        tiempos = []
        for _ in range(iteraciones):
            Gt = nx.gnm_random_graph(nodos, aristas, seed=rng.randint(0, 10 ** 9))
            for u, v in Gt.edges():
                Gt[u][v]["distancia"] = rng.uniform(1, 50)
            t0 = time.perf_counter()
            for o in Gt.nodes():
                dijkstra(Gt, o, "distancia")
            tiempos.append(time.perf_counter() - t0)
        filas.append((nodos, aristas, statistics.mean(tiempos), statistics.stdev(tiempos), min(tiempos), "aleatorio"))
    tiempos = []
    for _ in range(iteraciones):
        t0 = time.perf_counter()
        for o in G.nodes():
            dijkstra(G, o, "distancia")
        tiempos.append(time.perf_counter() - t0)
    filas.append((n_r, m_r, statistics.mean(tiempos), statistics.stdev(tiempos), min(tiempos), "REAL"))
    filas.sort(key=lambda f: (f[0], f[5] != "aleatorio"))

    print("\n" + "=" * 78)
    print("  PRUEBA DE TIEMPO DE EJECUCIÓN – TABLA 9 (3.6.3.c)")
    print(f"  Dijkstra desde todos los nodos | {iteraciones} repeticiones | semilla {semilla}")
    print("=" * 78)
    print(f"  {'Nodos':<8} {'Aristas':<9} {'Media (s)':>11} {'Desv. est.':>11} {'Mínimo (s)':>11}  Grafo")
    print("-" * 78)
    for n, m, me, sd, mn, tipo in filas:
        marca = "  ⭐ PROYECTO" if tipo == "REAL" else ""
        print(f"  {n:<8} {m:<9} {me:>11.5f} {sd:>11.5f} {mn:>11.5f}  {tipo}{marca}")
    print("=" * 78 + "\n")
    return filas


def probar_tiempo_ruta_dia(matrices, criterio="distancia"):
    print("\n" + "=" * 70)
    print("  PRUEBA DE TIEMPO - ORDEN ÓPTIMO DE LOS 7 DÍAS (TSP por permutaciones)")
    print("=" * 70)
    tiempos = []
    for dia in range(1, 8):
        destinos = [d for d in DIAS_CONFIG[dia]["destinos"] if d != BASE]
        t0 = time.perf_counter()
        ruta_optima_dia(destinos, matrices, criterio)
        dt = time.perf_counter() - t0
        tiempos.append(dt)
        print(f"  Día {dia} ({len(destinos)} destinos): {dt:.6f} s")
    print("-" * 50)
    print(f"  Promedio por día: {sum(tiempos) / len(tiempos):.6f} s | Total 7 días: {sum(tiempos):.6f} s")
    print("=" * 70 + "\n")
    return tiempos

# ═══════════════════════════════════════════════════════════════════════════════
# 7. EXPORTACIÓN A EXCEL
# ═══════════════════════════════════════════════════════════════════════════════

def exportar_excel(G, matrices, nombre="Rutas_Cocle.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    AZUL, GRIS, BLANCO = "1F4E79", "F2F2F2", "FFFFFF"

    def celda(ws, fila, col, valor, bold=False, bg=None, alinear="center", fmt=None):
        c = ws.cell(row=fila, column=col, value=valor)
        c.font = Font(name="Arial", bold=bold, size=10)
        if bg:
            c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal=alinear, vertical="center", wrap_text=True)
        if fmt:
            c.number_format = fmt
        return c

    def titulo_hoja(ws, texto, ncols):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        c = ws.cell(row=1, column=1, value=texto)
        c.font = Font(name="Arial", bold=True, color=BLANCO, size=14)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

    # Hoja 1: rutas
    ws = wb.create_sheet("Rutas_7_Dias")
    titulo_hoja(ws, "RUTAS TURÍSTICAS - PROVINCIA DE COCLÉ (criterio oficial: distancia)", 8)
    for j, h in enumerate(["Día", "Zona", "Ruta óptima (ida y vuelta)", "Distancia (km)", "Tiempo (min)",
                           "Costo (USD)", "Nivel", "Frase"], 1):
        celda(ws, 2, j, h, bold=True, bg=GRIS)
    for dia, cfg in DIAS_CONFIG.items():
        i = dia + 2
        bg = BLANCO if i % 2 == 0 else GRIS
        celda(ws, i, 1, dia, bg=bg)
        celda(ws, i, 2, cfg["zona"], bg=bg)
        celda(ws, i, 3, texto_ruta_dia(cfg), bg=bg)
        celda(ws, i, 4, cfg["distancia_km"], bg=bg, fmt="0.0")
        celda(ws, i, 5, cfg["tiempo_auto_min"], bg=bg, fmt="0")
        celda(ws, i, 6, cfg["costo_usd"], bg=bg, fmt="0.00")
        celda(ws, i, 7, cfg["nivel"], bg=bg)
        celda(ws, i, 8, cfg["frase"], bg=bg, alinear="left")
    for j, ancho in enumerate([6, 30, 42, 14, 12, 12, 10, 48], 1):
        ws.column_dimensions[get_column_letter(j)].width = ancho

    # Hoja 2: aristas (anexo de datos)
    ws2 = wb.create_sheet("Aristas")
    titulo_hoja(ws2, f"ARISTAS DEL GRAFO ({G.number_of_edges()})", 7)
    for j, h in enumerate(["Origen", "Destino", "Nombre origen", "Nombre destino", "Distancia (km)",
                           "Tiempo (min)", "Costo (USD)"], 1):
        celda(ws2, 2, j, h, bold=True, bg=GRIS)
    for k, (u, v, d) in enumerate(sorted(G.edges(data=True), key=lambda e: (min(e[0], e[1]), max(e[0], e[1]))), 3):
        bg = BLANCO if k % 2 == 1 else GRIS
        celda(ws2, k, 1, _cod(u), bg=bg)
        celda(ws2, k, 2, _cod(v), bg=bg)
        celda(ws2, k, 3, ATRACTIVOS[u]["nombre"], bg=bg, alinear="left")
        celda(ws2, k, 4, ATRACTIVOS[v]["nombre"], bg=bg, alinear="left")
        celda(ws2, k, 5, d["distancia"], bg=bg, fmt="0.00")
        celda(ws2, k, 6, d["tiempo"], bg=bg, fmt="0")
        celda(ws2, k, 7, round(d["costo"], 4), bg=bg, fmt="0.00")
    for j, ancho in enumerate([9, 9, 34, 34, 15, 13, 13], 1):
        ws2.column_dimensions[get_column_letter(j)].width = ancho

    # Hoja 3: matriz de distancias mínimas
    ws3 = wb.create_sheet("Matriz_Distancia")
    nodos = sorted(G.nodes())
    titulo_hoja(ws3, "MATRIZ DE DISTANCIAS MÍNIMAS (km) – Dijkstra", len(nodos) + 1)
    celda(ws3, 2, 1, "", bold=True, bg=GRIS)
    for j, n in enumerate(nodos, 2):
        celda(ws3, 2, j, _cod(n), bold=True, bg=GRIS)
    for i, o in enumerate(nodos, 3):
        celda(ws3, i, 1, _cod(o), bold=True, bg=GRIS)
        for j, d in enumerate(nodos, 2):
            celda(ws3, i, j, matrices["distancia"][o][d], fmt="0.0")
    for j in range(1, len(nodos) + 2):
        ws3.column_dimensions[get_column_letter(j)].width = 7

    wb.save(nombre)
    print(f"\n  ✔ Excel exportado: {nombre}")

# ═══════════════════════════════════════════════════════════════════════════════
# 8. CONSOLA Y MENÚS
# ═══════════════════════════════════════════════════════════════════════════════

def limpiar():
    os.system("cls" if os.name == "nt" else "clear")


def separador(char="─", ancho=64, color="\033[96m"):
    print(f"{color}{char * ancho}\033[0m")


def titulo(texto):
    limpiar()
    separador("═")
    print(f"\033[1;36m  {texto}\033[0m")
    separador("═")


def elegir_criterio():
    print("\n  Elige el criterio de optimización:")
    print("   [1] Distancia mínima (km)  ← criterio oficial")
    print("   [2] Tiempo mínimo (min)")
    print("   [3] Costo mínimo (USD)")
    while True:
        op = input("\n  Opción: ").strip()
        if op == "1":
            return "distancia"
        if op == "2":
            return "tiempo"
        if op == "3":
            return "costo"
        print("  ✗ Opción inválida.")


def mostrar_tabla_nodos(G):
    separador()
    print(f"\033[1m  {'ID':<4} {'Cód':<5} {'Nombre':<34} {'Tipo':<18} {'Distrito':<12}\033[0m")
    separador()
    for nid, data in sorted(ATRACTIVOS.items()):
        print(f"  {nid:<4} {data['cod']:<5} {data['nombre']:<34} {data['tipo']:<18} {data['distrito']:<12}")
    separador()


def mostrar_resultado_ruta(G, secuencia, matrices):
    mat_d, mat_t, mat_c = matrices["distancia"], matrices["tiempo"], matrices["costo"]
    separador()
    print(f"  \033[1m{'Tramo':<22} {'Distancia':>10} {'Tiempo':>10} {'Costo':>10}\033[0m")
    separador()
    tot_d = tot_t = tot_c = 0
    for i in range(len(secuencia) - 1):
        u, v = secuencia[i], secuencia[i + 1]
        d, t, c = mat_d[u][v] or 0, mat_t[u][v] or 0, mat_c[u][v] or 0
        tot_d += d
        tot_t += t
        tot_c += c
        print(f"  {_cod(u)} → {_cod(v):<16} {d:>9.1f}km {t:>9.1f}min {c:>9.2f}$")
    separador()
    print(f"  \033[1m{'TOTAL':<22} {tot_d:>9.1f}km {tot_t:>9.1f}min {tot_c:>9.2f}$\033[0m")
    separador()


def mostrar_todas_rutas():
    titulo("🗺️ RUTAS DE 7 DÍAS (criterio oficial: distancia, ida y vuelta a Penonomé)")
    for dia in range(1, 8):
        cfg = DIAS_CONFIG[dia]
        print(f"\n  📅 DÍA {dia}: {cfg['zona']}")
        print(f"  📍 {texto_ruta_dia(cfg)}")
        print(f"  📏 {cfg['distancia_km']:.1f} km | ⏱️ {cfg['tiempo_auto_min']} min | 💰 ${cfg['costo_usd']:.2f}")
        print(f"  💬 '{cfg['frase']}'")
        print("-" * 50)


def submenu_dia(G, matrices, caminos):
    titulo("RUTA ÓPTIMA POR DÍA")
    print("\n  Días disponibles:\n")
    for d, cfg in DIAS_CONFIG.items():
        print(f"   \033[1m[{d}]\033[0m  {cfg['zona']}")
        print(f"         Destinos: {' → '.join(_cod(n) for n in cfg['destinos'])}")
    print()
    while True:
        entrada = input("  ¿Qué día quieres ver? (1-7): ").strip()
        if entrada.isdigit() and 1 <= int(entrada) <= 7:
            dia = int(entrada)
            break
        print("  ✗ Ingresa un número del 1 al 7.")
    criterio = elegir_criterio()
    cfg = DIAS_CONFIG[dia]
    destinos = [d for d in cfg["destinos"] if d != BASE]
    orden, costo = ruta_optima_dia(destinos, matrices, criterio)
    if orden is None:
        print("  ✗ No se pudo calcular la ruta.")
        return
    secuencia = [BASE] + orden + [BASE]
    titulo(f"DÍA {dia} — {cfg['zona']}")
    print(f"\n  📝 {cfg['descripcion']}")
    print(f"  💬 '{cfg['frase']}'")
    print(f"\n  Criterio: {criterio.upper()} | Total: {costo:.1f}")
    print("\n  Orden óptimo de visita:")
    print(f"  \033[1;32m{' → '.join(_cod(n) for n in secuencia)}\033[0m\n")
    mostrar_resultado_ruta(G, secuencia, matrices)
    print("\n  \033[1m[1]\033[0m  Ver grafo de este día")
    print("  \033[1m[2]\033[0m  Guardar grafo como PNG")
    print("  \033[1m[3]\033[0m  Volver al menú")
    sub = input("\n  Opción: ").strip()
    if sub == "1":
        dibujar_grafo_dia(G, dia, matrices, caminos, criterio)
    elif sub == "2":
        dibujar_grafo_dia(G, dia, matrices, caminos, criterio, guardar=f"Grafo_Dia_{dia}_{criterio}.png")


def submenu_camino(G, matrices, caminos):
    titulo("CAMINO MÍNIMO ENTRE DOS NODOS")
    mostrar_tabla_nodos(G)
    print()
    while True:
        try:
            u = int(input("  ID del nodo ORIGEN (1-29): ").strip())
            v = int(input("  ID del nodo DESTINO (1-29): ").strip())
        except ValueError:
            print("  ✗ Ingresa un número válido.")
            continue
        if es_id_valido(u) and es_id_valido(v) and u != v:
            break
        print("  ✗ IDs inválidos o iguales.")
    criterio = elegir_criterio()
    camino = caminos[criterio][u][v]
    valor = matrices[criterio][u][v]
    if not camino:
        print(f"\n  ✗ No existe camino entre {u} y {v}.")
        return
    titulo(f"CAMINO MÍNIMO: {_cod(u)} → {_cod(v)}")
    print(f"\n  Criterio: {criterio.upper()} | Total: {valor:.2f}")
    print(f"  Ruta: {' → '.join(_cod(n) for n in camino)}\n")


def submenu_validaciones(G, matrices, caminos):
    titulo("VALIDACIONES (Sección 3.6.1)")
    print("""
  [1]  a) Generar plantilla CSV para validación externa
  [2]  a) Evaluar plantilla CSV ya completada
  [3]  b) Consistencia interna (subestructura óptima)
  [4]  c) Análisis de sensibilidad (distancia vs tiempo + Monte Carlo)
  [5]  Volver
""")
    op = input("  Opción: ").strip()
    if op == "1":
        generar_plantilla_validacion(G)
    elif op == "2":
        ruta = input("  Archivo CSV (Enter = validacion_externa_plantilla.csv): ").strip()
        evaluar_validacion_externa(G, ruta or "validacion_externa_plantilla.csv")
    elif op == "3":
        validar_consistencia_interna(G, matrices, caminos)
    elif op == "4":
        print("\n  ⏳ Ejecutando (unos segundos)...")
        analisis_sensibilidad(G, matrices, caminos)


def submenu_pruebas(G, matrices, caminos):
    titulo("PRUEBAS COMPUTACIONALES (Sección 3.6.3)")
    print("""
  [1]  b) Integridad de datos (automática)
  [2]  a) Caja negra automática + verificación contra networkx/Floyd-Warshall
  [3]  c) Tiempos de ejecución (Tabla 9)
  [4]  Volver
""")
    op = input("  Opción: ").strip()
    if op == "1":
        prueba_integridad_datos(G)
    elif op == "2":
        tabla_casos_caja_negra(G, matrices, caminos)
        verificacion_algoritmo(G, matrices, caminos)
    elif op == "3":
        print("\n  ⏳ Ejecutando pruebas...")
        probar_tiempo_ejecucion(G)
        probar_tiempo_ruta_dia(matrices)


def menu_principal():
    print("""
  \033[1;36m╔══════════════════════════════════════════════════════════════╗
  ║   RUTAS TURÍSTICAS – PROVINCIA DE COCLÉ, PANAMÁ            ║
  ║   Algoritmo de Dijkstra | Universidad de Panamá            ║
  ║   7 DÍAS · 29 NODOS · CRITERIO OFICIAL: DISTANCIA          ║
  ╚══════════════════════════════════════════════════════════════╝\033[0m

  \033[1m[1]\033[0m  Ver inventario de los 29 atractivos
  \033[1m[2]\033[0m  Ver grafo completo (nodos y aristas)
  \033[1m[3]\033[0m  Ver ruta óptima de un día (1–7)
  \033[1m[4]\033[0m  Ver todas las rutas de los 7 días
  \033[1m[5]\033[0m  Camino mínimo entre dos nodos (caja negra manual)
  \033[1m[6]\033[0m  Exportar todo a Excel
  \033[1m[7]\033[0m  Guardar grafos como imágenes PNG
  \033[1m[8]\033[0m  Validaciones (Sección 3.6.1)
  \033[1m[9]\033[0m  Pruebas computacionales (Sección 3.6.3)
  \033[1m[0]\033[0m  Salir
""")


def inicializar():
    G = construir_grafo()
    matrices, caminos = calcular_todas_matrices(G)
    recalcular_estadisticas_dias(matrices, criterio="distancia")
    return G, matrices, caminos


def ejecutar_todo_sin_menu():
    """python rutas_cocle_dijkstra.py --pruebas : genera todos los resultados."""
    G, matrices, caminos = inicializar()
    print(f"\n  Grafo: {G.number_of_nodes()} nodos · {G.number_of_edges()} aristas")
    mostrar_todas_rutas_sin_limpiar()
    prueba_integridad_datos(G)
    ok_alg, _ = verificacion_algoritmo(G, matrices, caminos)
    ok_cn = tabla_casos_caja_negra(G, matrices, caminos)
    ok_ci, _ = validar_consistencia_interna(G, matrices, caminos)
    analisis_sensibilidad(G, matrices, caminos)
    probar_tiempo_ejecucion(G)
    probar_tiempo_ruta_dia(matrices)
    exportar_excel(G, matrices)
    generar_plantilla_validacion(G)
    print("\n  RESUMEN: verificación algoritmo:", "OK" if ok_alg else "FALLA",
          "| caja negra:", "OK" if ok_cn else "FALLA",
          "| consistencia interna:", "OK" if ok_ci else "FALLA")


def mostrar_todas_rutas_sin_limpiar():
    print("\n  RUTAS DE 7 DÍAS (criterio oficial: distancia, ida y vuelta a Penonomé)")
    for dia in range(1, 8):
        cfg = DIAS_CONFIG[dia]
        print(f"  Día {dia}: {texto_ruta_dia(cfg)} | {cfg['distancia_km']:.1f} km | "
              f"{cfg['tiempo_auto_min']} min | ${cfg['costo_usd']:.2f}")


def main():
    limpiar()
    print("\n  \033[1;36mInicializando sistema...\033[0m")
    G, matrices, caminos = inicializar()
    print(f"  ✔ Grafo construido: {G.number_of_nodes()} nodos · {G.number_of_edges()} aristas")
    print("  ✔ Matrices calculadas para los 3 criterios.")
    print("  ✔ Estadísticas de los 7 días recalculadas (criterio oficial: distancia).")
    errores, advertencias = prueba_integridad_datos(G, verbose=False)
    if errores:
        print(f"  ❌ Integridad de datos: {len(errores)} ERRORES (ver menú 9 → 1)")
    else:
        print(f"  ✔ Integridad de datos: sin errores | {len(advertencias)} advertencias (ver menú 9 → 1)")
    input("\n  Presiona Enter para continuar...")

    try:
        while True:
            limpiar()
            menu_principal()
            opcion = input("  Opción: ").strip()
            if opcion == "0":
                limpiar()
                print("\n  ¡Hasta luego!\n")
                break
            elif opcion == "1":
                titulo("INVENTARIO DE ATRACTIVOS")
                mostrar_tabla_nodos(G)
            elif opcion == "2":
                titulo("GRAFO COMPLETO")
                dibujar_grafo_completo(G, elegir_criterio())
            elif opcion == "3":
                submenu_dia(G, matrices, caminos)
            elif opcion == "4":
                mostrar_todas_rutas()
            elif opcion == "5":
                print("\n  ℹ️ Prueba de Caja Negra manual (3.6.3.a): verifica a mano el camino mínimo.")
                submenu_camino(G, matrices, caminos)
            elif opcion == "6":
                titulo("EXPORTAR A EXCEL")
                try:
                    exportar_excel(G, matrices)
                except PermissionError:
                    print("\n  ✗ No se pudo escribir 'Rutas_Cocle.xlsx'. Ciérralo en Excel e inténtalo de nuevo.")
            elif opcion == "7":
                titulo("GUARDAR GRAFOS")
                criterio = elegir_criterio()
                os.makedirs("grafos_cocle", exist_ok=True)
                dibujar_grafo_completo(G, criterio, guardar=f"grafos_cocle/grafo_completo_{criterio}.png")
                for dia in range(1, 8):
                    dibujar_grafo_dia(G, dia, matrices, caminos, criterio,
                                      guardar=f"grafos_cocle/dia_{dia}_{criterio}.png")
                print("\n  ✔ Grafos guardados en 'grafos_cocle/'")
            elif opcion == "8":
                submenu_validaciones(G, matrices, caminos)
            elif opcion == "9":
                submenu_pruebas(G, matrices, caminos)
            else:
                print("  ✗ Opción no válida.")
            if opcion != "0":
                input("\n  Presiona Enter...")
    except (KeyboardInterrupt, EOFError):
        print("\n\n  Programa terminado.\n")


if __name__ == "__main__":
    if "--pruebas" in sys.argv:
        ejecutar_todo_sin_menu()
    else:
        main()
